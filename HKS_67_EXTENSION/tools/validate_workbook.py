"""Optional independent export check: pip install openpyxl; python tools/validate_workbook.py file.xlsx [expected_images]."""
import sys
from openpyxl import load_workbook
w = load_workbook(sys.argv[1])
assert w.sheetnames == ['Overview', 'Comments', 'Media', 'Coverage']
assert len(w['Overview']._charts) == 2
assert 'HKS1' in w['Comments'].tables or w['Comments'].max_row == 1
assert 'HKS2' in w['Media'].tables or w['Media'].max_row == 1
assert all(cell.data_type != 'f' for sheet in w for row in sheet for cell in row)
if w['Comments'].max_row > 1:
    assert w['Comments']['A2'].data_type == 's'
if len(sys.argv) > 2:
    assert len(w['Comments']._images) == int(sys.argv[2])
print({'sheets': w.sheetnames, 'comments': w['Comments'].max_row-1, 'media_records': w['Media'].max_row-1, 'charts': len(w['Overview']._charts), 'images': len(w['Comments']._images), 'formula_cells': 0})
